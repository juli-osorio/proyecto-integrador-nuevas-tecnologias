#DESARROLLO DEL TALLER

import openpyxl
from tabulate import tabulate
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# 1. Cargar el archivo Excel - abrirlo
wb= openpyxl.load_workbook('ventas_tecnologia_v2.xlsx', data_only=True) #valores y no formulas

#----------------------------

#2. Seleccionar la hoja activa
hoja= wb.active #la que estaba abierta cuando se guardó
print(f'Hoja cargada exitosamente: {hoja.title}')

#Lectura dinámica de registros
datos= []                #lista vacía para almacenar

for fila in hoja.iter_rows(min_row=2, values_only=True):  #hoja.iter_rows(): recorrer las filas // #Recorrer desde la fila 2 (omitir encabezados)
    datos.append(list(fila)) #list(fila) convierte esa tupla en una lista -> moificar

print(f'Total de registros leídos en la hoja: {len(datos)}')

#----------------------------

#3. Visualización Profesional en Consola (tabulate)
#importado al inicio: from tabulate import tabulate

#Nombre de columnas:
encabezados= ['ID', 'Producto', 'Categoría', 'Precio ($)', 'Ventas (Unid)']
#Alineación por columna:
alineacion= ('center', 'left', 'center', 'right', 'right')
#Mostrar tabla formateada en consola:
print('\n=== REPORTE DE VENTAS EN CONSOLA (22 PRODUCTOS) ===')
print(tabulate(datos, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))

#----------------------------

#4: Generación del Informe Ejecutivo en PDF (ReportLab)
#importado al inicio: 
#.  from reportlab.lib.pagesizes import letter
#.  from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
#.  from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
#.  from reportlab.lib import colors

# a. Crear documento base
doc = SimpleDocTemplate("reporte_ejecutivo_ventas.pdf", pagesize=letter)
estilos = getSampleStyleSheet()

# b. Estilos personalizados
estilo_titulo = ParagraphStyle(
    'EstiloTitulo', parent=estilos['Title'],
    fontName='Helvetica-Bold', fontSize=18,
    textColor=colors.HexColor("#1A365D"), spaceAfter=10
)
estilo_cuerpo = ParagraphStyle('EstiloCuerpo', parent=estilos['Normal'], fontSize=9)
estilo_encabezado = ParagraphStyle(
    'EstiloEncabezado', parent=estilos['Normal'],
    fontName='Helvetica-Bold', fontSize=10,
    textColor=colors.whitesmoke, alignment=1
)

# c. Lista story (carrito de compras)
story = [
    Paragraph("Informe Ejecutivo de Ventas Tecno", estilo_titulo),
    Paragraph(f"<b>Total de Productos Procesados:</b> {len(datos)}", estilo_cuerpo),
    Spacer(1, 12)
]

# d. Convertir celdas a Paragraph
encabezados_pdf = [Paragraph(f"<b>{h}</b>", estilo_encabezado) for h in encabezados]
filas_pdf = [
    [
        Paragraph(str(r[0]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
        Paragraph(str(r[1]), ParagraphStyle('L', parent=estilo_cuerpo, alignment=0)),
        Paragraph(str(r[2]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
        Paragraph(f"${r[3]:,.2f}", ParagraphStyle('R', parent=estilo_cuerpo, alignment=2)),
        Paragraph(str(r[4]), ParagraphStyle('R', parent=estilo_cuerpo, alignment=2))
    ]
    for r in datos
]

# e. Configurar anchos y tabla
anchos_cols = [40, 160, 100, 80, 80]
tabla_pdf = Table([encabezados_pdf] + filas_pdf, colWidths=anchos_cols, repeatRows=1)

# f. Aplicar estilo visual con TableStyle
tabla_pdf.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
]))

story.append(tabla_pdf)

# g. Compilar PDF
doc.build(story)

#-------------------------------------
#-------------------------------------
#-------------------------------------

#RETOS PRÁCTICOS
#-------------------------------------

#Reto 1 — Filtrado por Criterio de Precio (solo productos con precio > $100)

#Filtrar los datos
datos_filtrados = []
for fila in datos:
    if fila[3] is not None and fila[3] > 100:   # fila[3] = Precio
        datos_filtrados.append(fila)

#Mostrar en consola con tabulate
print(f'\n=== PRODUCTOS CON PRECIO MAYOR A $100 ({len(datos_filtrados)} PRODUCTOS) ===')
print(tabulate(datos_filtrados, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))

#Generar el PDF con los datos filtrados
doc_filtrado = SimpleDocTemplate("reporte_filtrado_precio.pdf", pagesize=letter)

story_filtrado = [
    Paragraph("Informe Ejecutivo de Ventas Tecno", estilo_titulo),
    Paragraph(f"<b>Productos con precio mayor a $100:</b> {len(datos_filtrados)}", estilo_cuerpo),
    Spacer(1, 12)
]

filas_filtradas_pdf = [
    [
        Paragraph(str(r[0]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
        Paragraph(str(r[1]), ParagraphStyle('L', parent=estilo_cuerpo, alignment=0)),
        Paragraph(str(r[2]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
        Paragraph(f"${r[3]:,.2f}", ParagraphStyle('R', parent=estilo_cuerpo, alignment=2)),
        Paragraph(str(r[4]), ParagraphStyle('R', parent=estilo_cuerpo, alignment=2))
    ]
    for r in datos_filtrados
]

tabla_filtrada = Table([encabezados_pdf] + filas_filtradas_pdf, colWidths=anchos_cols, repeatRows=1)
tabla_filtrada.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
]))

story_filtrado.append(tabla_filtrada)
doc_filtrado.build(story_filtrado)


#--------------------------------------------------------------------------------------

# RETO 2: Métricas y fila de resumen (sobre el catálogo completo)

#Calcular métricas
total_unidades = 0
suma_precios = 0
cantidad_productos = 0

for fila in datos:
    if fila[3] is not None and fila[4] is not None:
        total_unidades += fila[4]      # fila[4] = Ventas (Unid)
        suma_precios += fila[3]        # fila[3] = Precio
        cantidad_productos += 1

precio_promedio = suma_precios / cantidad_productos

#Mostrar métricas al final de la consola
print('\n=== MÉTRICAS DE RESUMEN ===')
print(f'Total de unidades vendidas: {total_unidades:,}')
print(f'Precio promedio de los productos: ${precio_promedio:,.2f}')

#Generar el PDF con la fila de resumen al final
doc_metricas = SimpleDocTemplate("reporte_metricas.pdf", pagesize=letter)

story_metricas = [
    Paragraph("Informe Ejecutivo de Ventas Tecno", estilo_titulo),
    Paragraph(f"<b>Total de Productos Procesados:</b> {len(datos)}", estilo_cuerpo),
    Spacer(1, 12)
]

#Fila de resumen (en negrita)
estilo_resumen = ParagraphStyle('Resumen', parent=estilo_cuerpo, fontName='Helvetica-Bold')
fila_resumen = [
    Paragraph("", estilo_resumen),
    Paragraph("TOTAL / PROMEDIO", estilo_resumen),
    Paragraph("", estilo_resumen),
    Paragraph(f"${precio_promedio:,.2f}", ParagraphStyle('RR', parent=estilo_resumen, alignment=2)),
    Paragraph(f"{total_unidades:,}", ParagraphStyle('RR', parent=estilo_resumen, alignment=2))
]

#Tabla: encabezados + filas del catálogo completo (filas_pdf ya existe) + fila resumen
tabla_metricas = Table([encabezados_pdf] + filas_pdf + [fila_resumen], colWidths=anchos_cols, repeatRows=1)
tabla_metricas.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),  # fondo gris en la última fila
]))

story_metricas.append(tabla_metricas)
doc_metricas.build(story_metricas)

#--------------------------------------------------------------------------------------


# RETO 3: Búsqueda dinámica por ID de producto

#Pedir el ID al usuario
id_buscado = input('\nIngresa el ID del producto que deseas buscar: ').strip()

#Buscar el ID en la matriz de datos
producto_encontrado = None
for fila in datos:
    if str(fila[0]) == id_buscado:     # fila[0] = ID
        producto_encontrado = fila
        break                          # ya lo encontró, deja de buscar

#Mostrar el resultado
if producto_encontrado:
    print(f'\n=== RESULTADO DE BÚSQUEDA: ID {id_buscado} ===')
    print(tabulate([producto_encontrado], headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))
else:
    print(f'\nADVERTENCIA: No existe ningún producto con el ID "{id_buscado}".')

#--------------------------------------------------------------------------------------


# RETO 4: Menú interactivo de consola con funciones

def ver_catalogo():
    print(f'\n=== CATÁLOGO COMPLETO ({len(datos)} PRODUCTOS) ===')
    print(tabulate(datos, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))


def filtrar_precio():
    filtrados = []
    for fila in datos:
        if fila[3] is not None and fila[3] > 100:
            filtrados.append(fila)
    print(f'\n=== PRODUCTOS CON PRECIO MAYOR A $100 ({len(filtrados)} PRODUCTOS) ===')
    print(tabulate(filtrados, headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))


def buscar_id():
    id_ingresado = input('\nIngresa el ID del producto: ').strip()
    for fila in datos:
        if str(fila[0]) == id_ingresado:
            print(f'\n=== RESULTADO DE BÚSQUEDA: ID {id_ingresado} ===')
            print(tabulate([fila], headers=encabezados, tablefmt="fancy_grid", colalign=alineacion))
            return                      # lo encontró, sale de la función
    print(f'\nADVERTENCIA: No existe ningún producto con el ID "{id_ingresado}".')


def exportar_pdf():
    documento = SimpleDocTemplate("reporte_ejecutivo_ventas.pdf", pagesize=letter)
    contenido = [
        Paragraph("Informe Ejecutivo de Ventas Tecno", estilo_titulo),
        Paragraph(f"<b>Total de Productos Procesados:</b> {len(datos)}", estilo_cuerpo),
        Spacer(1, 12)
    ]
    filas = [
        [
            Paragraph(str(r[0]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
            Paragraph(str(r[1]), ParagraphStyle('L', parent=estilo_cuerpo, alignment=0)),
            Paragraph(str(r[2]), ParagraphStyle('C', parent=estilo_cuerpo, alignment=1)),
            Paragraph(f"${r[3]:,.2f}", ParagraphStyle('R', parent=estilo_cuerpo, alignment=2)),
            Paragraph(str(r[4]), ParagraphStyle('R', parent=estilo_cuerpo, alignment=2))
        ]
        for r in datos
    ]
    encabezados_tabla = [Paragraph(f"<b>{h}</b>", estilo_encabezado) for h in encabezados]
    tabla = Table([encabezados_tabla] + filas, colWidths=anchos_cols, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A365D")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
    ]))
    contenido.append(tabla)
    documento.build(contenido)
    print("\n[ÉXITO] Archivo 'reporte_ejecutivo_ventas.pdf' generado correctamente.")


def menu():
    while True:
        print('\n========== MENÚ TECNOVENTAS ==========')
        print('1. Ver catálogo completo')
        print('2. Filtrar por precio (> $100)')
        print('3. Buscar por ID')
        print('4. Exportar informe PDF')
        print('5. Salir del programa')
        opcion = input('Elige una opción (1-5): ').strip()

        if opcion == '1':
            ver_catalogo()
        elif opcion == '2':
            filtrar_precio()
        elif opcion == '3':
            buscar_id()
        elif opcion == '4':
            exportar_pdf()
        elif opcion == '5':
            print('\n¡Hasta luego! Programa finalizado.')
            break      #(rompe el while True y termina)
        else:
            print('\nOpción no válida. Escribe un número del 1 al 5.')


#Iniciar el menú
menu()